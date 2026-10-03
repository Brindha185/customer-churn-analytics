"""
verify_excel_vs_sql.py
----------------------
Checks that the Excel workbook's formula results equal the SQL results (KPIs and several segment tables).
The workbook must have been calculated first (open + save in Excel, or LibreOffice recalc), because
openpyxl-written formulas have no cached values until then.

Usage: python tests/verify_excel_vs_sql.py <workbook.xlsx> <sql_results_dir>
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def main(xlsx, res):
    res = Path(res)
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Calculations"]
    xl = {ws[f"A{r}"].value: ws[f"B{r}"].value for r in range(4, 20)}
    if xl["Total Customers"] is None:
        sys.exit("Workbook has no calculated values yet. Open it in Excel, save, and re-run.")
    sql = pd.read_csv(res / "q00_kpi_summary.csv").iloc[0]
    pairs = [("Total Customers", "total_customers", 1), ("Active Customers", "active_customers", 1),
             ("Churned Customers", "churned_customers", 1), ("Overall Churn Rate", "churn_rate_pct", 100),
             ("Retention Rate", "retention_rate_pct", 100), ("Monthly Recurring Revenue (active)", "mrr_active", 1),
             ("Monthly Revenue Lost to Churn", "monthly_revenue_lost", 1),
             ("Average Monthly Charges", "avg_monthly_charges", 1),
             ("Average Customer Tenure (months)", "avg_tenure_months", 1),
             ("High-Risk Customers", "high_risk_customers", 1),
             ("Monthly Revenue at Risk", "monthly_revenue_at_risk", 1)]
    ok = True
    for xk, sk, m in pairs:
        a, b = xl[xk] * m, float(sql[sk])
        good = abs(a - b) < 0.01
        ok &= good
        print(f"{xk:<40}{a:>14.2f}{b:>14.2f}  {'OK' if good else 'MISMATCH'}")

    rows = list(wb["Pivot Analysis"].iter_rows(values_only=True))

    def block(title):
        for i, r in enumerate(rows):
            if r[0] == title:
                out, j = [], i + 2
                while j < len(rows) and rows[j][0] is not None:
                    out.append(rows[j])
                    j += 1
                return out
        raise KeyError(title)

    for title, qf, key in [("Churn by Contract", "q05_churn_by_contract", "contract"),
                           ("Churn by Tenure Group", "q08_churn_by_tenure_group", "tenure_group"),
                           ("Churn by Payment Method", "q07_churn_by_payment_method", "payment_method"),
                           ("Churn by Customer Value Segment (CLTV quartiles)", "q18_value_segment_churn", "value_segment"),
                           ("Churn by Monthly Charge Group", "q09_churn_by_monthly_charges", "charge_group")]:
        q = pd.read_csv(res / f"{qf}.csv").set_index(key)
        for r in block(title):
            s = q.loc[r[0]]
            good = r[1] == s.customers and r[2] == s.churned and abs(r[3] * 100 - s.churn_pct) < 0.01
            ok &= good
            if not good:
                print("MISMATCH", title, r, dict(s))
        print(f"{title}: checked")
    q14 = pd.read_csv(res / "q14_lifecycle_churn.csv")
    lb = block("Churn by Customer Lifecycle (months since signup)")
    for r, (_, s) in zip(lb, q14.iterrows()):
        good = r[3] == s.customers and r[4] == s.churned
        ok &= good
        if not good:
            print("MISMATCH lifecycle", r, dict(s))
    print("Lifecycle: checked", len(lb), "buckets")
    print("\nExcel vs SQL:", "PASS" if ok else "FAIL")
    return ok


if __name__ == "__main__":
    sys.exit(0 if main(sys.argv[1], sys.argv[2]) else 1)
