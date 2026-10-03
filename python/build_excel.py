"""
build_excel.py
--------------
Builds excel/customer_churn_analysis.xlsx from your raw + cleaned data.

Sheets:  Raw Data | Cleaned Data | Calculations | Pivot Analysis | Dashboard | Cleaning Log | Read Me

Everything analytical is a LIVE FORMULA (COUNTIFS, SUMIFS, AVERAGEIFS...), so you can inspect and
explain each number. Formulas are calculated when you open the file in Excel.

Note: Python cannot create real PivotTables or slicers. "Pivot Analysis" holds formula-driven
summary tables that behave like pivots; the 'Read Me' sheet gives the 6-step recipe for adding
real PivotTables + slicers on top of the Cleaned Data sheet (10 minutes, and good to show in interviews).
"""
import math
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
FONT = "Arial"
NAVY, GREY, RED, LIGHT = "1F3864", "595959", "C0392B", "F2F2F2"

DERIVED = ["tenure_group", "charge_group", "high_spender", "household_type", "value_segment",
           "contract_category", "services_count", "revenue_lost_monthly",
           "risk_profile_match", "high_risk_flag"]

DERIVED_NOTES = {
    "tenure_group": "Bands: 0-6, 7-12, 13-24, 25-48, 49-72 months (ASSUMPTION - adjust after viewing the distribution).",
    "charge_group": "Bands: under 35, 35-69.99, 70-89.99, 90+ (ASSUMPTION).",
    "high_spender": "1 if Monthly Charges >= 75th percentile (Calculations!B24).",
    "household_type": "Age substitute (dataset has no age): Senior / with dependents / with partner / single.",
    "value_segment": "CLTV quartiles: Low <= 25th percentile, High >= 75th percentile. CLTV is IBM-modelled.",
    "contract_category": "Month-to-month vs Long-term (1 or 2 year).",
    "services_count": "Number of add-on services = Yes (security, backup, device protection, tech support, streaming TV/movies).",
    "revenue_lost_monthly": "Monthly Charges if churned, else 0.",
    "risk_profile_match": "ASSUMPTION v1: Month-to-month AND tenure <= 12 AND has internet AND no Tech Support. Validate via Pivot Analysis rule back-test.",
    "high_risk_flag": "risk_profile_match AND customer is still active.",
}


def f(bold=False, size=10, color="000000", italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def header_style(cell):
    cell.font = f(True, 10, "FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def read_native(raw_path):
    p = Path(raw_path)
    return pd.read_excel(p) if p.suffix.lower() in (".xlsx", ".xls") else pd.read_csv(p)


def build(raw_path, clean_csv, log_csv, out_path):
    raw = read_native(raw_path)
    df = pd.read_csv(clean_csv, dtype={"zip_code": str})
    log = pd.read_csv(log_csv)
    n = len(df)
    last = n + 1

    wb = Workbook()
    ws_raw = wb.active
    ws_raw.title = "Raw Data"
    ws_cl = wb.create_sheet("Cleaned Data")
    ws_calc = wb.create_sheet("Calculations")
    ws_pv = wb.create_sheet("Pivot Analysis")
    ws_db = wb.create_sheet("Dashboard")
    ws_log = wb.create_sheet("Cleaning Log")
    ws_rm = wb.create_sheet("Read Me")

    # ------------------------------------------------------------------ Raw Data
    ws_raw.append(list(raw.columns))
    for row in raw.itertuples(index=False):
        ws_raw.append([None if (isinstance(v, float) and math.isnan(v)) else v for v in row])
    for c in ws_raw[1]:
        header_style(c)
    ws_raw.freeze_panes = "B2"
    ws_raw.sheet_properties.tabColor = "7F7F7F"

    # ------------------------------------------------------------------ Cleaned Data
    base_cols = list(df.columns)
    all_cols = base_cols + DERIVED
    L = {name: get_column_letter(i + 1) for i, name in enumerate(all_cols)}

    ws_cl.append(all_cols)
    for i, name in enumerate(all_cols, start=1):
        c = ws_cl.cell(row=1, column=i)
        header_style(c)
        if name in DERIVED:
            c.fill = PatternFill("solid", fgColor=RED)
            c.comment = Comment(DERIVED_NOTES[name], "analyst")

    for r, rec in enumerate(df.itertuples(index=False), start=2):
        for ci, v in enumerate(rec, start=1):
            if isinstance(v, float) and math.isnan(v):
                v = None
            ws_cl.cell(row=r, column=ci, value=v)
        g = lambda name: f"{L[name]}{r}"  # noqa: E731
        ws_cl[g("tenure_group")] = (f'=IF({g("tenure_months")}<=6,"0-6 mo",IF({g("tenure_months")}<=12,"7-12 mo",'
                                    f'IF({g("tenure_months")}<=24,"13-24 mo",IF({g("tenure_months")}<=48,"25-48 mo","49-72 mo"))))')
        ws_cl[g("charge_group")] = (f'=IF({g("monthly_charges")}<35,"Under 35",IF({g("monthly_charges")}<70,"35 to 69.99",'
                                    f'IF({g("monthly_charges")}<90,"70 to 89.99","90 and above")))')
        ws_cl[g("high_spender")] = f'=IF({g("monthly_charges")}>=Calculations!$B$24,1,0)'
        ws_cl[g("household_type")] = (f'=IF({g("senior_citizen")}="Yes","Senior",IF({g("dependents")}="Yes","Non-senior with dependents",'
                                      f'IF({g("partner")}="Yes","Non-senior with partner","Non-senior single")))')
        ws_cl[g("value_segment")] = (f'=IF({g("cltv")}<=Calculations!$B$22,"Low",IF({g("cltv")}>=Calculations!$B$23,"High","Mid"))')
        ws_cl[g("contract_category")] = f'=IF({g("contract")}="Month-to-month","Month-to-month","Long-term")'
        ws_cl[g("services_count")] = ("=" + "+".join(f'({g(a)}="Yes")' for a in
                                      ["online_security", "online_backup", "device_protection",
                                       "tech_support", "streaming_tv", "streaming_movies"]))
        ws_cl[g("revenue_lost_monthly")] = f'=IF({g("churn_flag")}=1,{g("monthly_charges")},0)'
        ws_cl[g("risk_profile_match")] = (f'=IF(AND({g("contract")}="Month-to-month",{g("tenure_months")}<=12,'
                                          f'{g("internet_service")}<>"No",{g("tech_support")}="No"),1,0)')
        ws_cl[g("high_risk_flag")] = f'=IF(AND({g("churn_flag")}=0,{g("risk_profile_match")}=1),1,0)'

    ws_cl.freeze_panes = "B2"
    ws_cl.auto_filter.ref = f"A1:{get_column_letter(len(all_cols))}{last}"
    for i, name in enumerate(all_cols, start=1):
        ws_cl.column_dimensions[get_column_letter(i)].width = max(12, min(22, len(name) + 3))
    ws_cl.sheet_properties.tabColor = NAVY

    def R(name):
        """Absolute range for a Cleaned Data column."""
        return f"'Cleaned Data'!${L[name]}$2:${L[name]}${last}"

    # ------------------------------------------------------------------ Calculations
    ws_calc["A1"] = "KPI Calculations"
    ws_calc["A1"].font = f(True, 14, NAVY)
    ws_calc["A2"] = ("All values are formulas on the Cleaned Data sheet. Retention/churn here = churned customers "
                     "divided by all customers in the file (single snapshot, not a month-over-month rate).")
    ws_calc["A2"].font = f(False, 9, GREY, True)
    for c, t in zip("ABC", ["KPI", "Value", "Formula logic / meaning"]):
        ws_calc[f"{c}3"] = t
        header_style(ws_calc[f"{c}3"])

    kpis = [
        (4, "Total Customers", f"=COUNTA({R('customer_id')})", "#,##0", "Count of CustomerID"),
        (5, "Active Customers", f"=COUNTIF({R('churn_flag')},0)", "#,##0", "Customers with churn_flag = 0"),
        (6, "Churned Customers", f"=COUNTIF({R('churn_flag')},1)", "#,##0", "Customers with churn_flag = 1"),
        (7, "Overall Churn Rate", "=B6/B4", "0.0%", "Churned / Total"),
        (8, "Retention Rate", "=1-B7", "0.0%", "1 - Churn Rate"),
        (9, "Monthly Recurring Revenue (active)", f"=SUMIFS({R('monthly_charges')},{R('churn_flag')},0)", "$#,##0", "Sum of Monthly Charges, active customers"),
        (10, "Monthly Revenue Lost to Churn", f"=SUM({R('revenue_lost_monthly')})", "$#,##0", "Sum of Monthly Charges, churned customers"),
        (11, "Annualised Revenue Lost (ASSUMPTION)", "=B10*12", "$#,##0", "Monthly lost x 12; assumes customers would have stayed a full year"),
        (12, "Average Monthly Charges", f"=AVERAGE({R('monthly_charges')})", "$#,##0.00", "All customers"),
        (13, "Average Customer Tenure (months)", f"=AVERAGE({R('tenure_months')})", "0.0", "All customers"),
        (14, "Avg Tenure - Active (months)", f"=AVERAGEIFS({R('tenure_months')},{R('churn_flag')},0)", "0.0", "Active customers"),
        (15, "Avg Tenure - Churned (months)", f"=AVERAGEIFS({R('tenure_months')},{R('churn_flag')},1)", "0.0", "Churned customers"),
        (16, "High-Risk Customers", f"=SUM({R('high_risk_flag')})", "#,##0", "Active customers matching the risk profile (assumption v1)"),
        (17, "Monthly Revenue at Risk", f"=SUMIFS({R('monthly_charges')},{R('high_risk_flag')},1)", "$#,##0", "Monthly Charges of high-risk customers"),
        (18, "Risk Profile Precision", f"=IFERROR(SUMIFS({R('churn_flag')},{R('risk_profile_match')},1)/SUM({R('risk_profile_match')}),0)", "0.0%", "Churn rate among ALL customers matching the profile (back-test)"),
        (19, "Risk Profile Capture", f"=IFERROR(SUMIFS({R('churn_flag')},{R('risk_profile_match')},1)/B6,0)", "0.0%", "Share of all churners that match the profile (back-test)"),
    ]
    for row, label, formula, fmt, note in kpis:
        ws_calc[f"A{row}"] = label
        ws_calc[f"B{row}"] = formula
        ws_calc[f"B{row}"].number_format = fmt
        ws_calc[f"C{row}"] = note
        for c in "ABC":
            ws_calc[f"{c}{row}"].font = f(c == "B", 10, "000000" if c != "C" else GREY)

    ws_calc["A21"] = "Percentile cut-offs used by derived columns"
    ws_calc["A21"].font = f(True, 11, NAVY)
    params = [
        (22, "CLTV 25th percentile", f"=PERCENTILE({R('cltv')},0.25)", "Low value segment: CLTV <= this"),
        (23, "CLTV 75th percentile", f"=PERCENTILE({R('cltv')},0.75)", "High value segment: CLTV >= this"),
        (24, "Monthly Charges 75th percentile", f"=PERCENTILE({R('monthly_charges')},0.75)", "High spender: Monthly Charges >= this"),
    ]
    for row, label, formula, note in params:
        ws_calc[f"A{row}"] = label
        ws_calc[f"B{row}"] = formula
        ws_calc[f"B{row}"].number_format = "#,##0.00"
        ws_calc[f"C{row}"] = note
        for c in "ABC":
            ws_calc[f"{c}{row}"].font = f(c == "B", 10, "000000" if c != "C" else GREY)
    ws_calc.column_dimensions["A"].width = 40
    ws_calc.column_dimensions["B"].width = 16
    ws_calc.column_dimensions["C"].width = 75
    ws_calc.sheet_properties.tabColor = "2E75B6"

    # ------------------------------------------------------------------ Pivot Analysis
    # helper: label ordering computed once in pandas (labels only; all numbers are formulas)
    def rate_order(col, min_n=0, top=None):
        g = df.groupby(col).agg(n=("churn_flag", "size"), r=("churn_flag", "mean"))
        g = g[g.n >= min_n].sort_values("r", ascending=False)
        return list(g.index[:top] if top else g.index)

    ws_pv["A1"] = "Pivot Analysis - formula-driven summary tables"
    ws_pv["A1"].font = f(True, 14, NAVY)
    ws_pv["A2"] = ("Each table: Customers, Churned, Churn Rate, Monthly Revenue Lost, % of All Churn. Always read the rate "
                   "together with the customer count. Real PivotTables + slicers: see the Read Me sheet.")
    ws_pv["A2"].font = f(False, 9, GREY, True)

    blocks = {}
    state = {"row": 4}
    ex = {"churn": R("churn_flag"), "rev": R("revenue_lost_monthly")}

    def seg_block(key, title, dim, labels, note=None):
        r0 = state["row"]
        ws_pv[f"A{r0}"] = title
        ws_pv[f"A{r0}"].font = f(True, 11, NAVY)
        if note:
            ws_pv[f"C{r0}"] = note
            ws_pv[f"C{r0}"].font = f(False, 9, GREY, True)
        heads = ["Segment", "Customers", "Churned", "Churn Rate", "Monthly Revenue Lost", "% of All Churn"]
        for j, h in enumerate(heads, start=1):
            c = ws_pv.cell(row=r0 + 1, column=j, value=h)
            header_style(c)
        first = r0 + 2
        for k, lab in enumerate(labels):
            r = first + k
            ws_pv[f"A{r}"] = lab
            ws_pv[f"B{r}"] = f"=COUNTIFS({R(dim)},$A{r})"
            ws_pv[f"C{r}"] = f"=COUNTIFS({R(dim)},$A{r},{ex['churn']},1)"
            ws_pv[f"D{r}"] = f"=IFERROR(C{r}/B{r},0)"
            ws_pv[f"E{r}"] = f"=SUMIFS({ex['rev']},{R(dim)},$A{r})"
            ws_pv[f"F{r}"] = f"=IFERROR(C{r}/Calculations!$B$6,0)"
            ws_pv[f"B{r}"].number_format = "#,##0"
            ws_pv[f"C{r}"].number_format = "#,##0"
            ws_pv[f"D{r}"].number_format = "0.0%"
            ws_pv[f"E{r}"].number_format = "$#,##0"
            ws_pv[f"F{r}"].number_format = "0.0%"
            for c in "ABCDEF":
                ws_pv[f"{c}{r}"].font = f()
        lastr = first + len(labels) - 1
        ws_pv.conditional_formatting.add(
            f"D{first}:D{lastr}",
            ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F4B183"))
        blocks[key] = (first, lastr)
        state["row"] = lastr + 3

    seg_block("contract", "Churn by Contract", "contract", rate_order("contract"))
    seg_block("internet", "Churn by Internet Service", "internet_service", rate_order("internet_service"))
    seg_block("payment", "Churn by Payment Method", "payment_method", rate_order("payment_method"))
    seg_block("tenure", "Churn by Tenure Group", "tenure_group",
              ["0-6 mo", "7-12 mo", "13-24 mo", "25-48 mo", "49-72 mo"])
    seg_block("charge", "Churn by Monthly Charge Group", "charge_group",
              ["Under 35", "35 to 69.99", "70 to 89.99", "90 and above"])
    seg_block("household", "Churn by Household Type (age substitute)", "household_type",
              ["Senior", "Non-senior with dependents", "Non-senior with partner", "Non-senior single"])
    seg_block("value", "Churn by Customer Value Segment (CLTV quartiles)", "value_segment", ["Low", "Mid", "High"])
    seg_block("services", "Churn by Number of Add-on Services (plan proxy)", "services_count", list(range(0, 7)))
    seg_block("techsup", "Churn by Tech Support (support proxy)", "tech_support", ["Yes", "No"])
    seg_block("gender", "Churn by Gender", "gender", sorted(df["gender"].unique()))
    seg_block("paperless", "Churn by Paperless Billing", "paperless_billing", sorted(df["paperless_billing"].unique()))
    city_labels = rate_order("city", min_n=30, top=15)
    if city_labels:
        seg_block("city", "Churn by City - top 15 by churn rate, minimum 30 customers", "city", city_labels,
                  "Region substitute: dataset is California only. Threshold of 30 is an assumption.")

    # Lifecycle (tenure buckets of 6 months)
    r0 = state["row"]
    ws_pv[f"A{r0}"] = "Churn by Customer Lifecycle (months since signup)"
    ws_pv[f"A{r0}"].font = f(True, 11, NAVY)
    ws_pv[f"C{r0}"] = "No date column in this dataset: this is a lifecycle view, NOT a calendar trend."
    ws_pv[f"C{r0}"].font = f(False, 9, RED, True)
    for j, h in enumerate(["Months", "From", "To", "Customers", "Churned", "Churn Rate"], start=1):
        header_style(ws_pv.cell(row=r0 + 1, column=j, value=h))
    n_buck = max(1, math.ceil(df["tenure_months"].max() / 6))
    first = r0 + 2
    for k in range(n_buck):
        r = first + k
        lo, hi = (0 if k == 0 else k * 6 + 1), (k + 1) * 6
        ws_pv[f"A{r}"] = f"{lo} to {hi}"
        ws_pv[f"B{r}"] = lo
        ws_pv[f"C{r}"] = hi
        ws_pv[f"D{r}"] = f'=COUNTIFS({R("tenure_months")},">="&B{r},{R("tenure_months")},"<="&C{r})'
        ws_pv[f"E{r}"] = f'=COUNTIFS({R("tenure_months")},">="&B{r},{R("tenure_months")},"<="&C{r},{ex["churn"]},1)'
        ws_pv[f"F{r}"] = f"=IFERROR(E{r}/D{r},0)"
        ws_pv[f"F{r}"].number_format = "0.0%"
        for c in "ABCDEF":
            ws_pv[f"{c}{r}"].font = f()
    lastr = first + n_buck - 1
    ws_pv.conditional_formatting.add(
        f"F{first}:F{lastr}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F4B183"))
    blocks["lifecycle"] = (first, lastr)
    state["row"] = lastr + 3

    # Churn reasons
    reasons = (df.loc[df.churn_flag == 1, "churn_reason"].value_counts().head(10).index.tolist())
    if reasons:
        r0 = state["row"]
        ws_pv[f"A{r0}"] = "Top Stated Churn Reasons (churned customers only; self-reported)"
        ws_pv[f"A{r0}"].font = f(True, 11, NAVY)
        for j, h in enumerate(["Churn Reason", "Churners", "% of Churners", "Monthly Revenue Lost"], start=1):
            header_style(ws_pv.cell(row=r0 + 1, column=j, value=h))
        first = r0 + 2
        for k, lab in enumerate(reasons):
            r = first + k
            ws_pv[f"A{r}"] = lab
            ws_pv[f"B{r}"] = f"=COUNTIFS({R('churn_reason')},$A{r},{ex['churn']},1)"
            ws_pv[f"C{r}"] = f"=IFERROR(B{r}/Calculations!$B$6,0)"
            ws_pv[f"D{r}"] = f"=SUMIFS({ex['rev']},{R('churn_reason')},$A{r})"
            ws_pv[f"C{r}"].number_format = "0.0%"
            ws_pv[f"D{r}"].number_format = "$#,##0"
            for c in "ABCD":
                ws_pv[f"{c}{r}"].font = f()
        blocks["reasons"] = (first, first + len(reasons) - 1)
        state["row"] = first + len(reasons) + 2

    # Matrix: contract x tenure group
    r0 = state["row"]
    ws_pv[f"A{r0}"] = "Segment Comparison: Churn Rate by Contract x Tenure Group"
    ws_pv[f"A{r0}"].font = f(True, 11, NAVY)
    tgs = ["0-6 mo", "7-12 mo", "13-24 mo", "25-48 mo", "49-72 mo"]
    ws_pv.cell(row=r0 + 1, column=1, value="Contract")
    header_style(ws_pv.cell(row=r0 + 1, column=1))
    for j, t in enumerate(tgs, start=2):
        header_style(ws_pv.cell(row=r0 + 1, column=j, value=t))
    for k, ctr in enumerate(["Month-to-month", "One year", "Two year"]):
        r = r0 + 2 + k
        ws_pv[f"A{r}"] = ctr
        ws_pv[f"A{r}"].font = f()
        for j in range(len(tgs)):
            col = get_column_letter(j + 2)
            hdr = f"{col}${r0 + 1}"
            ws_pv[f"{col}{r}"] = (f'=IFERROR(COUNTIFS({R("contract")},$A{r},{R("tenure_group")},{hdr},{ex["churn"]},1)/'
                                  f'COUNTIFS({R("contract")},$A{r},{R("tenure_group")},{hdr}),"")')
            ws_pv[f"{col}{r}"].number_format = "0.0%"
            ws_pv[f"{col}{r}"].font = f()
    ws_pv.conditional_formatting.add(
        f"B{r0 + 2}:F{r0 + 4}",
        ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="E06666"))
    ws_pv[f"A{r0 + 5}"] = "Blank cell = no customers in that combination."
    ws_pv[f"A{r0 + 5}"].font = f(False, 9, GREY, True)
    state["row"] = r0 + 8

    # Rule back-test
    r0 = state["row"]
    ws_pv[f"A{r0}"] = "Risk Profile Back-test (is the high-risk rule any good?)"
    ws_pv[f"A{r0}"].font = f(True, 11, NAVY)
    for j, h in enumerate(["Group", "Customers", "Churned", "Churn Rate", "% of All Churners", "Lift vs Overall"], start=1):
        header_style(ws_pv.cell(row=r0 + 1, column=j, value=h))
    for k, (lab, val) in enumerate([("Matches risk profile", 1), ("Does not match", 0)]):
        r = r0 + 2 + k
        ws_pv[f"A{r}"] = lab
        ws_pv[f"B{r}"] = f"=COUNTIFS({R('risk_profile_match')},{val})"
        ws_pv[f"C{r}"] = f"=COUNTIFS({R('risk_profile_match')},{val},{ex['churn']},1)"
        ws_pv[f"D{r}"] = f"=IFERROR(C{r}/B{r},0)"
        ws_pv[f"E{r}"] = f"=IFERROR(C{r}/Calculations!$B$6,0)"
        ws_pv[f"F{r}"] = f"=IFERROR(D{r}/Calculations!$B$7,0)"
        for c, fmt in zip("BCDEF", ["#,##0", "#,##0", "0.0%", "0.0%", "0.00"]):
            ws_pv[f"{c}{r}"].number_format = fmt
        for c in "ABCDEF":
            ws_pv[f"{c}{r}"].font = f()
    ws_pv.column_dimensions["A"].width = 38
    for c in "BCDEF":
        ws_pv.column_dimensions[c].width = 18
    ws_pv.sheet_properties.tabColor = "2E75B6"

    # ------------------------------------------------------------------ Dashboard
    ws_db.sheet_view.showGridLines = False
    for i in range(1, 11):
        ws_db.column_dimensions[get_column_letter(i)].width = 14
    ws_db.merge_cells("A1:J1")
    ws_db["A1"] = "Customer Churn & Retention Dashboard"
    ws_db["A1"].font = f(True, 18, NAVY)
    ws_db.merge_cells("A2:J2")
    ws_db["A2"] = "Source: IBM Telco Customer Churn sample dataset (fictional). Single snapshot, no calendar dates."
    ws_db["A2"].font = f(False, 9, GREY, True)

    cards = [
        ("Total Customers", "Calculations!B4", "#,##0"), ("Active Customers", "Calculations!B5", "#,##0"),
        ("Churned Customers", "Calculations!B6", "#,##0"), ("Churn Rate", "Calculations!B7", "0.0%"),
        ("Retention Rate", "Calculations!B8", "0.0%"),
        ("Monthly Recurring Revenue", "Calculations!B9", "$#,##0"), ("Monthly Revenue Lost", "Calculations!B10", "$#,##0"),
        ("Avg Monthly Charges", "Calculations!B12", "$#,##0.00"), ("Avg Tenure (months)", "Calculations!B13", "0.0"),
        ("High-Risk Customers", "Calculations!B16", "#,##0"),
    ]
    thin = Side(style="thin", color="D9D9D9")
    for idx, (label, ref, fmt) in enumerate(cards):
        row = 4 if idx < 5 else 7
        c1 = (idx % 5) * 2 + 1
        a, b = get_column_letter(c1), get_column_letter(c1 + 1)
        ws_db.merge_cells(f"{a}{row}:{b}{row}")
        ws_db.merge_cells(f"{a}{row + 1}:{b}{row + 1}")
        ws_db[f"{a}{row}"] = label
        ws_db[f"{a}{row}"].font = f(False, 9, GREY)
        ws_db[f"{a}{row + 1}"] = f"={ref}"
        accent = RED if label in ("Churn Rate", "Monthly Revenue Lost", "High-Risk Customers") else NAVY
        ws_db[f"{a}{row + 1}"].font = f(True, 18, accent)
        ws_db[f"{a}{row + 1}"].number_format = fmt
        for rr in (row, row + 1):
            for cc in (a, b):
                cell = ws_db[f"{cc}{rr}"]
                cell.fill = PatternFill("solid", fgColor=LIGHT)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = Border(top=thin if rr == row else None, bottom=thin if rr == row + 1 else None,
                                     left=thin if cc == a else None, right=thin if cc == b else None)
        ws_db.row_dimensions[row + 1].height = 30

    def bar(key, title, anchor, col_cat=1, col_val=4, horizontal=False, line=False):
        first, lastr = blocks[key]
        ch = LineChart() if line else BarChart()
        if not line:
            ch.type = "bar" if horizontal else "col"
        ch.title = title
        ch.legend = None
        ch.height, ch.width = 7.5, 15.5
        data = Reference(ws_pv, min_col=col_val, min_row=first, max_row=lastr)
        cats = Reference(ws_pv, min_col=col_cat, min_row=first, max_row=lastr)
        ch.add_data(data, titles_from_data=False)
        ch.set_categories(cats)
        ch.y_axis.number_format = "0%"
        ch.y_axis.majorGridlines = None
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        if horizontal:
            ch.x_axis.scaling.orientation = "maxMin"  # highest category on top
        s = ch.series[0]
        if line:
            s.smooth = False
        s.graphicalProperties.solidFill = RED if key in ("contract", "lifecycle") else "2E75B6"
        s.graphicalProperties.line.solidFill = RED if key in ("contract", "lifecycle") else "2E75B6"
        ws_db.add_chart(ch, anchor)

    bar("contract", "Churn rate by contract", "A10", horizontal=True)
    bar("lifecycle", "Churn rate by months since signup (lifecycle view)", "F10", col_val=6, line=True)
    bar("tenure", "Churn rate by tenure group", "A26")
    bar("payment", "Churn rate by payment method", "F26", horizontal=True)
    bar("internet", "Churn rate by internet service", "A42")
    bar("value", "Churn rate by customer value (CLTV)", "F42")

    ws_db.sheet_properties.tabColor = RED
    wb.active = wb.sheetnames.index("Dashboard")

    # ------------------------------------------------------------------ Cleaning Log
    ws_log.append(["Check", "Rows affected", "Action taken"])
    for c in ws_log[1]:
        header_style(c)
    for rec in log.itertuples(index=False):
        ws_log.append(list(rec))
    ws_log.column_dimensions["A"].width = 70
    ws_log.column_dimensions["B"].width = 14
    ws_log.column_dimensions["C"].width = 85
    for row in ws_log.iter_rows(min_row=2):
        for c in row:
            c.font = f()
    ws_log.sheet_properties.tabColor = "7F7F7F"

    # ------------------------------------------------------------------ Read Me
    lines = [
        ("Customer Churn & Retention Analytics - Excel workbook", "title"),
        ("", None),
        ("Sheets", "h"),
        ("Raw Data: the file exactly as downloaded. Never edited.", None),
        ("Cleaned Data: cleaned table + 10 derived columns (red headers; hover a header for its rule). Derived columns are formulas.", None),
        ("Calculations: KPI formulas and the percentile cut-offs used by the derived columns.", None),
        ("Pivot Analysis: formula-driven summary tables (COUNTIFS/SUMIFS) behind the dashboard charts.", None),
        ("Dashboard: KPI cards + charts, all linked to the sheets above.", None),
        ("Cleaning Log: every check performed, rows affected, action taken.", None),
        ("", None),
        ("Assumptions (also in docs/assumptions.md)", "h"),
        ("Total Charges blank -> 0 (customer not yet billed, normally tenure = 0).", None),
        ("Tenure bands 0-6 / 7-12 / 13-24 / 25-48 / 49-72 months; charge bands under 35 / 35-69.99 / 70-89.99 / 90+.", None),
        ("Value segment = CLTV quartiles (CLTV is modelled by IBM, not observed). High spender = top quarter of Monthly Charges.", None),
        ("High-risk profile v1: Month-to-month AND tenure <= 12 AND has internet AND no Tech Support. Back-test it in Pivot Analysis; revise if it performs poorly.", None),
        ("Annualised revenue lost = monthly x 12 (assumes the customer would have stayed a full year).", None),
        ("City table needs >= 30 customers per city. No dates in the data, so the 'trend' chart is a lifecycle view.", None),
        ("", None),
        ("Add real PivotTables + slicers (about 10 minutes)", "h"),
        ("1. Cleaned Data sheet: click any cell in the data -> Insert -> Table (name it tblChurn).", None),
        ("2. Insert -> PivotTable -> From Table -> New Worksheet.", None),
        ("3. Rows: contract. Values: Count of customer_id, and Average of churn_flag (format as %). Average of churn_flag IS the churn rate.", None),
        ("4. Insert -> PivotChart (clustered bar). Repeat for tenure_group, payment_method, internet_service, city.", None),
        ("5. Insert -> Slicer on contract, internet_service, senior_citizen, tenure_group.", None),
        ("6. Right-click a slicer -> Report Connections -> tick every pivot so one click filters them all.", None),
        ("", None),
        ("Formulas are calculated when the file opens in Excel. If values look blank, press F9 or enable calculation.", None),
    ]
    for i, (text, kind) in enumerate(lines, start=1):
        c = ws_rm.cell(row=i, column=1, value=text)
        c.font = (f(True, 14, NAVY) if kind == "title" else f(True, 11, NAVY) if kind == "h" else f())
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws_rm.column_dimensions["A"].width = 130
    ws_rm.sheet_properties.tabColor = "7F7F7F"

    wb.calculation.fullCalcOnLoad = True
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Excel workbook written: {out_path}")
    return out_path


if __name__ == "__main__":
    raw = sys.argv[1] if len(sys.argv) > 1 else ROOT / "data" / "raw" / "Telco_customer_churn.xlsx"
    build(raw, ROOT / "data" / "cleaned" / "churn_clean.csv", ROOT / "data" / "cleaned" / "cleaning_log.csv",
          ROOT / "excel" / "customer_churn_analysis.xlsx")
